"""콘텐츠 → Word(.docx). 수정용 원고: PDF와 같은 내용, Word 기본 기능만 사용."""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from textbook.blocks import blanks, tokens
from textbook.meta import BOOK, PARTS
from textbook.render_html import mock_label, unit_tests

FONT = "맑은 고딕"
INK, SUB, MUTE = "1E2433", "5B6475", "8A93A6"
RED, GREEN, LINE = "D64545", "1F9254", "DCE1E9"
TAG_COLOR = {"S": "2B59C3", "V": "C93A5A", "O": "0E8A7E", "I.O": "0E8A7E", "D.O": "2C9C6A",
             "C": "D9701A", "O.C": "D9701A", "M": "8A93A6"}
CIRCLED = "①②③④⑤"
MODE = {"teacher": False, "xref": {}}
FIG_PATHS = {}  # build.py가 rasterize_figs 결과를 채움


def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(k), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = rgb(color)


def run_shade(run, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    run._element.get_or_add_rPr().append(shd)


def cell_shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill.lstrip("#"))
    tcpr.append(shd)


def cell_borders(cell, **edges):
    """edges: top/bottom/left/right = (size_eighths, color) 또는 None(없음)"""
    tcpr = cell._element.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        v = edges.get(edge)
        if v:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(v[0]))
            e.set(qn("w:color"), v[1].lstrip("#"))
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tcpr.append(b)


def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcpr = cell._element.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcpr.append(m)


def para_border_bottom(p, color=LINE, sz=6):
    ppr = p._element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:color"), color)
    bdr.append(e)
    ppr.append(bdr)


def para_border_left(p, color, sz=24):
    ppr = p._element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    e = OxmlElement("w:left")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "6")
    e.set(qn("w:color"), color)
    bdr.append(e)
    ppr.append(bdr)


def spacing(p, before=0, after=4, line=1.25):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_inline(p, text, color, size=9.5, base_bold=False, base_color=INK):
    for style, s in tokens(str(text)):
        if style == "blank":
            if MODE["teacher"]:
                r = p.add_run(f" {s} ")
                set_font(r, size, True, RED)
            else:
                r = p.add_run("\u00a0" * max(6, int(len(s) * 2.6)))
                set_font(r, size)
            r.font.underline = True
            continue
        if style == "tag":
            r = p.add_run(f" {s} ")
            set_font(r, size - 2, True, "FFFFFF")
            run_shade(r, TAG_COLOR.get(s, MUTE))
            continue
        parts = s.split("\n")
        for i, seg in enumerate(parts):
            if i:
                p.add_run().add_break()
            r = p.add_run(seg)
            set_font(r, size, base_bold, base_color)
            if style == "b":
                set_font(r, bold=True, color=color)
            elif style == "hl":
                r.font.underline = True
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif style == "x":
                r.font.strike = True
                set_font(r, color=RED)
            elif style == "o":
                set_font(r, bold=True, color=GREEN)


class Writer:
    def __init__(self):
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(9.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Mm(18))
        sec.top_margin, sec.bottom_margin = Mm(20), Mm(18)
        self.color = INK
        self.tint = "F5F7FA"

    # ---------- 기본 ----------
    def para(self, text="", size=9.5, bold=False, color=INK, align=None, before=0, after=4):
        p = self.doc.add_paragraph()
        spacing(p, before, after)
        if text:
            add_inline(p, text, self.color, size, bold, color)
        if align:
            p.alignment = align
        return p

    def table(self, rows, cols, widths=None):
        t = self.doc.add_table(rows=rows, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if widths:
            total = sum(widths)
            for i, w in enumerate(widths):
                for c in t.columns[i].cells:
                    c.width = Mm(174 * w / total)
        return t

    def cell_text(self, cell, text, size=9, bold=False, color=INK, align=None):
        p = cell.paragraphs[0]
        spacing(p, 0, 0, 1.2)
        add_inline(p, text, self.color, size, bold, color)
        if align:
            p.alignment = align
        return p

    def gap(self, pt=4):
        p = self.doc.add_paragraph()
        spacing(p, 0, 0, 1)
        p.paragraph_format.space_after = Pt(pt)
        r = p.add_run()
        r.font.size = Pt(2)

    def box(self, label, title, lines, fill, border, badge):
        t = self.table(1, 1, [1])
        c = t.cell(0, 0)
        cell_shade(c, fill)
        edge = (6, border)
        cell_borders(c, top=edge, bottom=edge, left=(24, badge), right=edge)
        cell_margins(c, 100, 100, 160, 160)
        p = c.paragraphs[0]
        spacing(p, 0, 2)
        r = p.add_run(f" {label} ")
        set_font(r, 7.5, True, "FFFFFF")
        run_shade(r, badge)
        r = p.add_run("  ")
        add_inline(p, title, self.color, 9.5, True)
        for line in lines:
            q = c.add_paragraph()
            spacing(q, 0, 1)
            q.paragraph_format.left_indent = Mm(3)
            add_inline(q, "• " + line if isinstance(line, str) else line, self.color, 9)
        self.gap(6)

    # ---------- 블록 ----------
    def h(self, text):
        p = self.para(text, 13, True, INK, before=12, after=6)
        para_border_left(p, self.color)
        p.paragraph_format.keep_with_next = True
        p.style = self.doc.styles["Heading 2"]
        for r in p.runs:
            set_font(r, 13, True, INK)
        return p

    def tbl(self, b):
        if b["caption"]:
            p = self.para("■ " + b["caption"], 8.5, True, self.color, after=2)
            p.paragraph_format.keep_with_next = True
        t = self.table(len(b["rows"]) + 1, len(b["head"]), b["widths"] or [1] * len(b["head"]))
        for i, h in enumerate(b["head"]):
            c = t.cell(0, i)
            cell_shade(c, self.tint)
            cell_borders(c, top=(16, self.color), bottom=(8, self.color))
            self.cell_text(c, h, 9, True, align=WD_ALIGN_PARAGRAPH.CENTER if i in b["center"] else None)
        for ri, row in enumerate(b["rows"], 1):
            last = ri == len(b["rows"])
            for ci, v in enumerate(row):
                c = t.cell(ri, ci)
                cell_borders(c, bottom=(12, INK) if last else (4, LINE))
                self.cell_text(c, v, 9, align=WD_ALIGN_PARAGRAPH.CENTER if ci in b["center"] else None)
        self.gap(6)

    def ex(self, b):
        for i, (en, ko) in enumerate(b["items"], 1):
            p = self.para(after=0)
            r = p.add_run(f"{i}  ")
            set_font(r, 8, True, self.color)
            add_inline(p, en, self.color, 10)
            p.paragraph_format.keep_with_next = True
            q = self.para(ko, 8.5, color=SUB, after=3)
            q.paragraph_format.left_indent = Mm(4.5)
        self.gap(4)

    def parse(self, b):
        segs = b["segs"]
        t = self.table(2, len(segs), [max(len(s), 4) for s, _ in segs])
        for i, (text, label) in enumerate(segs):
            col = TAG_COLOR.get(label, "FFFFFF") if label else "FFFFFF"
            c1, c2 = t.cell(0, i), t.cell(1, i)
            for c in (c1, c2):
                cell_shade(c, "F5F7FA")
            cell_borders(c1, bottom=(18, col) if label else None)
            cell_borders(c2)
            self.cell_text(c1, text, 10, True, INK if label else SUB, WD_ALIGN_PARAGRAPH.CENTER)
            self.cell_text(c2, label or "", 7.5, True, TAG_COLOR.get(label, MUTE), WD_ALIGN_PARAGRAPH.CENTER)
        if b["ko"]:
            self.para(b["ko"], 8.5, color=SUB, after=6, before=2)
        else:
            self.gap(6)

    def formula(self, b):
        t = self.table(1, 1, [1])
        c = t.cell(0, 0)
        e = (12, self.color)
        cell_borders(c, top=e, bottom=e, left=e, right=e)
        cell_margins(c, 120, 120, 160, 160)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(" 공식 ")
        set_font(r, 7.5, True, "FFFFFF")
        run_shade(r, self.color)
        r = p.add_run("  ")
        add_inline(p, b["text"], self.color, 11, True)
        if b["sub"]:
            q = c.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(q, b["sub"], self.color, 9, False, SUB)
        self.gap(6)

    def warn(self, b):
        t = self.table(len(b["items"]) * 2 + 1, 2, [1, 1])
        head = t.cell(0, 0).merge(t.cell(0, 1))
        for row in t.rows:
            for c in row.cells:
                cell_shade(c, "FFF5F5")
                cell_borders(c)
        p = head.paragraphs[0]
        r = p.add_run(" 주의 ")
        set_font(r, 7.5, True, "FFFFFF")
        run_shade(r, RED)
        add_inline(p, "  " + b["title"], self.color, 9.5, True)
        for i, (w, ok, why) in enumerate(b["items"]):
            c1, c2 = t.cell(1 + i * 2, 0), t.cell(1 + i * 2, 1)
            p = c1.paragraphs[0]
            add_inline(p, "✗ ", self.color, 9, True, RED)
            add_inline(p, w, self.color, 9)
            p = c2.paragraphs[0]
            add_inline(p, "✓ ", self.color, 9, True, GREEN)
            add_inline(p, ok, self.color, 9)
            why_c = t.cell(2 + i * 2, 0).merge(t.cell(2 + i * 2, 1))
            self.cell_text(why_c, "→ " + why, 8.5, color=SUB)
        self.gap(6)

    def check(self, b, unit_no):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        t = self.table(1, 1, [1])
        c = t.cell(0, 0)
        cell_shade(c, self.color)
        cell_borders(c)
        cell_margins(c, 100, 100, 180, 180)
        p = c.paragraphs[0]
        r = p.add_run("CHECK-UP")
        set_font(r, 14, True, "FFFFFF")
        r = p.add_run(f"   Unit {unit_no:02d} 확인 문제")
        set_font(r, 9, False, "FFFFFF")
        n = 1
        for step, desc, qs in b["steps"]:
            p = self.para(before=10, after=4)
            r = p.add_run(f" {step} ")
            set_font(r, 8, True, "FFFFFF")
            run_shade(r, self.color)
            r = p.add_run("  " + desc)
            set_font(r, 10, True, INK)
            for q in qs:
                p = self.para(after=2, before=4)
                r = p.add_run(f"{n:02d}  ")
                set_font(r, 10, True, self.color)
                add_inline(p, q["prompt"], self.color, 9.5)
                p.paragraph_format.keep_with_next = bool(q["options"] or q["lines"])
                if q["options"]:
                    for i, o in enumerate(q["options"]):
                        op = self.para(after=0)
                        op.paragraph_format.left_indent = Mm(7)
                        r = op.add_run(CIRCLED[i] + " ")
                        set_font(r, 9.5, True, self.color)
                        add_inline(op, o, self.color, 9.5)
                for _ in range(q["lines"]):
                    lp = self.para(after=0)
                    lp.paragraph_format.left_indent = Mm(7)
                    lp.paragraph_format.space_before = Pt(10)
                    para_border_bottom(lp, "B9C1CF")
                last = self.para(after=4)
                para_border_bottom(last)
                n += 1

    def chain(self, b):
        t = self.table(1, 1, [1])
        c = t.cell(0, 0)
        e = (6, self.color)
        cell_borders(c, top=e, bottom=e, left=e, right=e)
        cell_margins(c, 100, 100, 160, 160)
        p = c.paragraphs[0]
        add_inline(p, "↻ " + b["title"], self.color, 9.5, True, self.color)
        for i, (label, sent, note) in enumerate(b["steps"]):
            if i:
                q = c.add_paragraph()
                spacing(q, 0, 0)
                q.paragraph_format.left_indent = Mm(6)
                add_inline(q, "↓  " + (note or ""), self.color, 8, False, SUB)
            q = c.add_paragraph()
            spacing(q, 1, 1)
            r = q.add_run(f" {label} ")
            set_font(r, 8, True, self.color)
            run_shade(r, self.tint)
            q.add_run("  ")
            add_inline(q, sent, self.color, 10)
        self.gap(6)

    def anchor(self, b):
        t = self.table(1, 1, [1])
        c = t.cell(0, 0)
        cell_shade(c, self.tint)
        e = (18, self.color)
        cell_borders(c, top=e, bottom=e, left=e, right=e)
        cell_margins(c, 140, 140, 200, 200)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, "이 단원의 기준 문장", self.color, 8, True, self.color)
        for txt, size, col, bold in ((b["s"], 15, INK, True), (b["ko"], 9, SUB, False), (b.get("note") or "", 8, MUTE, False)):
            if not txt:
                continue
            q = c.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            spacing(q, 2, 0)
            add_inline(q, txt, self.color, size, bold, col)
        self.gap(8)

    def link(self, b):
        p = self.para(after=6)
        r = p.add_run(f" ↔ Unit {b['unit']:02d} ")
        set_font(r, 7.5, True, "FFFFFF")
        run_shade(r, SUB)
        add_inline(p, "  " + b["text"], self.color, 8.5, False, SUB)

    def stepup(self, b):
        p = self.para(before=14, after=6)
        para_border_bottom(p, LINE, 8)
        r = p.add_run(" STEP UP ")
        set_font(r, 9, True, "FFFFFF")
        run_shade(r, INK)
        r = p.add_run("  심화   ")
        set_font(r, 10, True, INK)
        add_inline(p, b["text"], self.color, 8.5, False, MUTE)

    def test(self, test, label, sub, newpage=True):
        if newpage:
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        p = self.para(after=2)
        r = p.add_run(f" {label} ")
        set_font(r, 9, True, "FFFFFF")
        run_shade(r, self.color)
        r = p.add_run(f"   {sub}")
        set_font(r, 9, False, SUB)
        total = sum(it["points"] for it in test["items"])
        p = self.para(test["title"], 16, True, after=2)
        p = self.para(f"이름 ____________    점수 ______ / {total}", 9, True, after=8)
        para_border_bottom(p, INK, 18)
        for n, it in enumerate(test["items"], 1):
            p = self.para(before=8, after=2)
            r = p.add_run(f"{n}. ")
            set_font(r, 10.5, True, self.color)
            if it["k"] == "wr":
                r = p.add_run("[서술형] ")
                set_font(r, 9, True, self.color)
            add_inline(p, it["prompt"], self.color, 9.5)
            r = p.add_run(f'  [{it["level"]}]' + (" [고난도]" if it["hard"] else "") + f' ({it["points"]}점)')
            set_font(r, 8, False, MUTE)
            p.paragraph_format.keep_with_next = True
            if it["passage"]:
                t = self.table(1, 1, [1])
                c = t.cell(0, 0)
                cell_shade(c, "FBFBFC")
                e = (4, LINE)
                cell_borders(c, top=e, bottom=e, left=e, right=e)
                cell_margins(c, 100, 100, 140, 140)
                add_inline(c.paragraphs[0], it["passage"], self.color, 9)
            if it["k"] == "mc":
                for i, o in enumerate(it["options"]):
                    op = self.para(after=0)
                    op.paragraph_format.left_indent = Mm(6)
                    r = op.add_run("①②③④⑤"[i] + " ")
                    set_font(r, 9.5, True, self.color)
                    add_inline(op, o, self.color, 9.5)
            else:
                if it["conditions"]:
                    cp = self.para(after=0)
                    cp.paragraph_format.left_indent = Mm(6)
                    add_inline(cp, "<조건> " + " / ".join(it["conditions"]), self.color, 8.5, False, SUB)
                for _ in range(it["lines"]):
                    lp = self.para(after=0)
                    lp.paragraph_format.left_indent = Mm(6)
                    lp.paragraph_format.space_before = Pt(10)
                    para_border_bottom(lp, "B9C1CF")
            if MODE["teacher"]:
                ans = "①②③④⑤"[it["answer"] - 1] if it["k"] == "mc" else it["answer"]
                ap = self.para(after=2)
                ap.paragraph_format.left_indent = Mm(6)
                add_inline(ap, f"정답 {ans}  ", self.color, 8.5, True, RED)
                add_inline(ap, it["explain"], self.color, 8.5, False, SUB)

    # ---------- 섹션 ----------
    def new_section(self, head_left, head_right, color):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        self._header(sec, head_left, head_right, color)
        return sec

    def _header(self, sec, left, right, color):
        hp = sec.header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.tab_stops.add_tab_stop(Mm(174), alignment=2)
        r = hp.add_run(left)
        set_font(r, 7.5, False, MUTE)
        r = hp.add_run("\t" + right)
        set_font(r, 7.5, True, color)
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run()
        set_font(r, 9, True, INK)
        for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), tag)
                r._element.append(fc)
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = txt
                r._element.append(it)


def render_blocks(w, blocks, unit_no):
    d = w.doc
    for b in blocks:
        t = b["t"]
        if t == "h":
            w.h(b["text"])
        elif t == "p":
            w.para(b["text"], after=6)
        elif t == "table":
            w.tbl(b)
        elif t == "ex":
            w.ex(b)
        elif t == "parse":
            w.parse(b)
        elif t == "formula":
            w.formula(b)
        elif t == "tip":
            w.box("TIP", b["title"], b["lines"], "FFF9E6", "F2DC9B", "E2A400")
        elif t in ("note", "list"):
            w.box("보충" if t == "note" else "정리", b["title"] or "", b.get("lines") or b.get("items"),
                  "F5F7FA", LINE, SUB)
        elif t == "warn":
            w.warn(b)
        elif t == "check":
            w.check(b, unit_no)
        elif t == "fig":
            from textbook.render_pdf import fig_key
            path = FIG_PATHS.get(fig_key(b["html"], w.color))
            if path and path.exists():
                p = w.para(after=2)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(path), width=Mm(b["width"] or 160))
            if b["caption"]:
                w.para(b["caption"], 8.5, color=SUB, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif t == "tnote":
            if MODE["teacher"]:
                w.box("교사 메모", b["title"], b["lines"], "FFF5F5", "F4CACA", RED)
        elif t == "anchor":
            w.anchor(b)
        elif t == "pilgi":
            w.box("✓", b["title"], b["lines"], "FFFDF6", "E6DFC8", w.color)
        elif t == "quick":
            w.box("▶", b["title"], b["lines"], w.tint, w.color, w.color)
        elif t == "chain":
            w.chain(b)
        elif t == "link":
            w.link(b)
        elif t == "stepup":
            w.stepup(b)
        elif t == "pb":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_docx(units, out_path: Path, appendix=None, part_tests=None, mocks=None, with_answers=True):
    part_tests, mocks = part_tests or {}, mocks or []
    w = Writer()
    d = w.doc

    # 표지
    w.para(BOOK["series"], 10, True, "2B59C3", after=80)
    for tag in ("S", "V", "O", "C", "M"):
        pass
    p = w.para(after=6)
    for tag in ("S", "V", "O", "C", "M"):
        r = p.add_run(f" {tag} ")
        set_font(r, 11, True, "FFFFFF")
        run_shade(r, TAG_COLOR[tag])
        p.add_run(" ")
    w.para(BOOK["title"], 44, True, INK, after=2)
    w.para(BOOK["en"], 16, False, MUTE, after=10)
    w.para(BOOK["sub"], 12, True, SUB, after=6)
    w.para((MODE.get("book", "") + " · " if MODE.get("book") else "") + ("교사용 · 정답 포함" if MODE["teacher"] else "학생용"),
           10, True, INK, after=80)
    t = w.table(len(PARTS), 3, [18, 28, 54])
    for i, (pid, (pno, pname, pen, color, tint)) in enumerate(PARTS.items()):
        names = " · ".join(u["title"] for u in units if u["part"] == pid)
        cells = t.rows[i].cells
        for c in cells:
            cell_shade(c, tint)
            cell_borders(c, bottom=(12, "FFFFFF"))
        cell_borders(cells[0], left=(36, color), bottom=(12, "FFFFFF"))
        w.cell_text(cells[0], pno, 9, True, color)
        w.cell_text(cells[1], pname, 11, True)
        w.cell_text(cells[2], names, 8.5, color=SUB)

    # 목차
    w.new_section(BOOK["series"], "목차", INK)
    w.para("CONTENTS", 9, True, "2B59C3", after=0)
    p = w.para("목차", 24, True, after=12)
    para_border_bottom(p, INK, 18)
    for pid, (pno, pname, pen, color, tint) in PARTS.items():
        us = [u for u in units if u["part"] == pid]
        if not us:
            continue
        w.color = color
        p = w.para(f"**{pno}**   {pname}   ", 11, True, before=8, after=3)
        r = p.add_run(pen)
        set_font(r, 8.5, False, MUTE)
        para_border_left(p, color)
        for u in us:
            p = w.para(after=1)
            p.paragraph_format.left_indent = Mm(4)
            r = p.add_run(f'{u["no"]:02d}   ')
            set_font(r, 10, True, color)
            r = p.add_run(u["title"])
            set_font(r, 10.5, True, INK)
            r = p.add_run(f'   {u["en"]}   ' + "★" * u["level"] + "☆" * (3 - u["level"]))
            set_font(r, 8, False, MUTE)
    w.color = INK
    w.para("※ 쪽 번호는 PDF판 기준입니다. Word에서는 왼쪽 '탐색 창'(보기 → 탐색 창)으로 단원을 이동하세요.", 8.5,
           color=SUB, before=14)

    # 단원
    for u in units:
        pno, pname, pen, color, tint = PARTS[u["part"]]
        w.color, w.tint = color, tint
        w.new_section(f'{BOOK["series"]} · {pno} {pname}', f'Unit {u["no"]:02d}  {u["title"]}', color)
        p = w.para(after=4)
        r = p.add_run(f" {pno} · {pname} ")
        set_font(r, 8, True, color)
        run_shade(r, tint)
        r = p.add_run("     난이도 " + "★" * u["level"] + "☆" * (3 - u["level"]))
        set_font(r, 8, True, color)
        p = w.para(after=8)
        r = p.add_run(f'UNIT {u["no"]:02d}  ')
        set_font(r, 26, True, color)
        r = p.add_run(u["title"])
        set_font(r, 22, True, INK)
        p.style = d.styles["Heading 1"]
        for r in p.runs:
            r.font.color.rgb = rgb(color if r.text.startswith("UNIT") else INK)
        r = p.add_run("\n" + u["en"])
        set_font(r, 10, False, MUTE)
        if MODE["xref"].get(u["key"]):
            r = p.add_run(f'    ▶ 문제편 p.{MODE["xref"][u["key"]]}')
            set_font(r, 10, True, color)
        para_border_bottom(p, color, 20)
        if u["intro"]:
            w.para(u["intro"], 9.5, color=SUB, before=6, after=6)
        if u["goals"]:
            w.box("목표", "학습 목표", [f"{i}. {g}" for i, g in enumerate(u["goals"], 1)], tint, tint, color)

        render_blocks(w, u["blocks"], u["no"])
        for i, (name, t) in enumerate(unit_tests(u)):
            w.test(t, name, f'Unit {u["no"]:02d} {u["title"]}', newpage=bool(u["blocks"]) or i > 0)
        nxt = units[units.index(u) + 1] if units.index(u) + 1 < len(units) else None
        if (nxt is None or nxt["part"] != u["part"]) and u["part"] in part_tests:
            w.test(part_tests[u["part"]], f"{pno} 종합 TEST", pname)

    w.color, w.tint = INK, "F0F2F6"
    for k, t in enumerate(mocks, 1):
        w.new_section(BOOK["series"], mock_label(k), INK)
        w.test(t, mock_label(k), "전 범위", newpage=False)

    # 부록
    if appendix:
        w.color, w.tint = INK, "F0F2F6"
        w.new_section(BOOK["series"], "부록", INK)
        w.para("APPENDIX", 9, True, "2B59C3", after=0)
        p = w.para("부록", 24, True, after=10)
        para_border_bottom(p, INK, 18)
        render_blocks(w, appendix, 0)

    # 정답 (학생용은 별책으로 분리)
    if not with_answers:
        d.core_properties.title = BOOK["title"]
        d.save(out_path)
        return
    w.color = INK
    w.new_section(BOOK["series"], "정답 및 해설", INK)
    w.para("ANSWERS", 9, True, "2B59C3", after=0)
    p = w.para("정답 및 해설", 24, True, after=10)
    para_border_bottom(p, INK, 18)
    for u in units:
        color = PARTS[u["part"]][3]
        w.color = color
        p = w.para(f'Unit {u["no"]:02d}  {u["title"]}', 11, True, color, before=10, after=4)
        para_border_left(p, color)
        bl = [x for b in u["blocks"] if b["t"] == "pilgi" for l in b["lines"] for x in blanks(l)]
        if bl:
            w.para("필기 포인트 빈칸", 8.5, True, color, before=4, after=1)
            w.para("  ·  ".join(f"{i}) {x}" for i, x in enumerate(bl, 1)), 8.5, color=SUB, after=2)
        n = 1
        for b in u["blocks"]:
            if b["t"] != "check":
                continue
            for step, desc, qs in b["steps"]:
                w.para(f"{step} {desc}", 8.5, True, color, before=4, after=1)
                for q in qs:
                    p = w.para(after=1)
                    r = p.add_run(f"{n:02d}  ")
                    set_font(r, 9, True, color)
                    add_inline(p, q["answer"], color, 9, True)
                    if q["explain"]:
                        add_inline(p, "   " + q["explain"], color, 8.5, False, SUB)
                    n += 1

    d.core_properties.title = BOOK["title"]
    d.core_properties.author = BOOK["series"]
    d.save(out_path)
